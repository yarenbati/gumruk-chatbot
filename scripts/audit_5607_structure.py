"""Offline structural audit for the supplied 5607 normalized source."""

from __future__ import annotations

import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

import docx
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "raw" / "5607-kacakcilikla-mucadele-kanunu.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
ARTICLE_RE = re.compile(r"^MADDE\s+(\d+)(?:\s*/\s*([A-Z]))?\s*[–—-]\s*(.*)$", re.I)
EK_RE = re.compile(r"^EK\s+MADDE\s+(\d+)(?:\s*/\s*([A-Z]))?\s*[–—-]?\s*(.*)$", re.I)
GECICI_RE = re.compile(r"^GEÇİCİ\s+MADDE\s+(\d+)(?:\s*/\s*([A-Z]))?\s*[–—-]?\s*(.*)$", re.I)
SECTION_RE = re.compile(r"^(BİRİNCİ|İKİNCİ|ÜÇÜNCÜ|DÖRDÜNCÜ|BEŞİNCİ|ALTINCI|YEDİNCİ|SEKİZİNCİ|DOKUZUNCU|ONUNCU|ONBİRİNCİ|ONİKİNCİ|ONÜÇÜNCÜ)\s+(KISIM|BÖLÜM|AYIRIM)$", re.I)
NUMBERED_RE = re.compile(r"^(?:\(\d+\)|\d+\.)\s+")
LETTERED_RE = re.compile(r"^[a-zçğıöşü]\)\s+", re.I)


def _article_match(text: str) -> tuple[str, str] | None:
    for kind, pattern in (("ek", EK_RE), ("gecici", GECICI_RE), ("normal", ARTICLE_RE)):
        match = pattern.match(text.strip())
        if match:
            no = match.group(1) + (f"/{match.group(2)}" if match.group(2) else "")
            return kind, no
    return None


def _footnotes(document: docx.document.Document) -> tuple[int, int, list[str]]:
    refs = sum(len(p._p.findall(".//w:footnoteReference", NS)) for p in document.paragraphs)
    parts = [p for p in document.part.package.parts if p.partname == "/word/footnotes.xml"]
    if not parts:
        return refs, 0, []
    root = etree.fromstring(parts[0].blob)
    bodies: list[str] = []
    for node in root.findall("w:footnote", NS):
        if node.get(f"{{{W_NS}}}type") in {"separator", "continuationSeparator"}:
            continue
        bodies.append("".join(node.itertext()).strip())
    return refs, len(bodies), bodies


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    document = docx.Document(str(SOURCE))
    texts = [p.text.strip() for p in document.paragraphs]
    nonempty = [t for t in texts if t]
    starts: list[dict[str, object]] = []
    for index, text in enumerate(texts):
        matched = _article_match(text)
        if matched:
            starts.append({"paragraph_index": index, "article_type": matched[0], "article_no": matched[1], "text": text})
    sections = [t for t in nonempty if SECTION_RE.fullmatch(t)]
    section_counts = Counter(SECTION_RE.fullmatch(t).group(2).lower() for t in sections)
    refs, bodies, body_samples = _footnotes(document)
    result = {
        "source": str(SOURCE.relative_to(ROOT)),
        "total_word_paragraphs": len(texts),
        "nonempty_paragraphs": len(nonempty),
        "tables": [{"rows": len(t.rows), "columns": len(t.columns)} for t in document.tables],
        "section_counts": dict(section_counts),
        "section_headings": sections,
        "article_starts": starts,
        "article_counts": dict(Counter(str(x["article_type"]) for x in starts)),
        "numbered_fikra_paragraphs": sum(bool(NUMBERED_RE.match(t)) for t in nonempty),
        "lettered_bent_paragraphs": sum(bool(LETTERED_RE.match(t)) for t in nonempty),
        "mulga_occurrences": sum(t.lower().count("mülga") for t in nonempty),
        "footnote_reference_occurrences": refs,
        "footnote_body_count": bodies,
        "footnote_body_samples": body_samples[:3],
        "replacement_characters": sum(t.count("\ufffd") for t in nonempty),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
