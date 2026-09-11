import io
import zipfile

from docx import Document

from src.annex_ingest import AnnexSourceFile, detect_file_type, ingest_member
from src.source_identity import AnnexSourceKey


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("first")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    doc.add_paragraph("last")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _source(name: str, data: bytes, detected: str) -> AnnexSourceFile:
    return AnnexSourceFile("gumruk_yonetmeligi", "test.zip", "z", name, name, "m", detected, ".docx")


def test_docx_paragraph_and_table_order_is_preserved() -> None:
    data = _docx_bytes()
    units = ingest_member(data, "EK-7.docx", source=_source("EK-7.docx", data, "OOXML_DOCX"))
    assert [b.kind for b in units[0].blocks] == ["paragraph", "table", "paragraph"]
    assert units[0].blocks[1].rows == (("left", "right"),)


def test_xlsx_signature_and_rtf_mismatch_route_correctly() -> None:
    xlsx = io.BytesIO()
    with zipfile.ZipFile(xlsx, "w") as z:
        z.writestr("xl/workbook.xml", '<workbook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Sheet One" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", '<worksheet><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>value</t></is></c><c r="B1"><f>SUM(A1)</f><v>2</v></c></row></sheetData></worksheet>')
    assert detect_file_type(xlsx.getvalue()) == "OOXML_XLSX"
    xlsx_units = ingest_member(xlsx.getvalue(), "EK-83.xlsx", source=_source("EK-83.xlsx", xlsx.getvalue(), "OOXML_XLSX"))
    spreadsheet = xlsx_units[0].blocks[0]
    assert spreadsheet.sheet_name == "Sheet One"
    assert spreadsheet.cells == (("A1", "value"), ("B1", {"value": "2", "formula": "SUM(A1)", "cached_value": "2"}))
    rtf = b"{\\rtf1\\ansi EK-48\\par Unicode: \\u304?}"
    assert detect_file_type(rtf) == "RTF"
    units = ingest_member(rtf, "EK-48.doc", source=_source("EK-48.doc", rtf, "RTF"))
    assert units[0].source_key == AnnexSourceKey.from_label("EK-48")
    assert units[0].source_member.detected_file_type == "RTF"
    assert units[0].warnings == ()


def test_heading_splits_one_physical_docx_without_splitting_inline_reference() -> None:
    doc = Document()
    doc.add_paragraph("EK-77")
    doc.add_paragraph("Body refers to EK-80 inline.")
    doc.add_paragraph("EK-77/A")
    doc.add_paragraph("Sub body")
    stream = io.BytesIO(); doc.save(stream); data = stream.getvalue()
    units = ingest_member(data, "EK-77.docx", source=_source("EK-77.docx", data, "OOXML_DOCX"))
    assert [u.human_label for u in units] == ["EK-77", "EK-77/A"]
    assert "EK-80" in units[0].rendered_text
